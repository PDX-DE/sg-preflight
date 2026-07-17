"""Size-analysis workbook auto-generation.

When `workbook_finder.resolve_workbook` returns `unavailable` for a profile but
the BMW pipeline has produced raw export-size data locally (CSV or JSON), this
module renders a Format A workbook to `~/sgfx_outputs/<profile>/delivery-workbook/<PROFILE>_auto_<YYYYMMDD>.xlsx`
matching the canonical project workbook schema audit.

The generated workbook is classified `auto_generated_locally` (NOT `from_ci`)
in the resulting `WorkbookCandidate` so the delivery-checklist surface can be
honest about provenance per `[[phase-j-automated-verdict-trajectory]]`.

Schema enforced (Format A — date-stamped IDC_EVO):
- Row 1: `[<PROFILE>, "", "", "", "", "", ""]`
- Row 2: `["Variant", "TextureCube", "Texture2D", "ArrayResource", "Effect", "Total", "Valeo est."]`
- Row 3+: one row per powertrain variant
  - Col F (Total) = sum of B-E
  - Col G (Valeo est.) = Total * 0.88

Raw data discovery patterns:
- `<bmw_root>/cars/BMW/<PROFILE>_EVO/export/size_data.{csv,json}`
- `<bmw_root>/cars/BMW/<PROFILE>/export/size_data.{csv,json}`
- `<workspace>/Cars/BMW/<PROFILE>/export/size_data.{csv,json}`
- `<workspace>/Cars_IDCevo/BMW/<PROFILE>/size_data.{csv,json}`

If none of the above resolve, the generator returns None and the delivery-checklist
surface falls back to the existing honest unavailable wording.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import shutil
import subprocess
from typing import Iterable

from sg_preflight.bmw_delivery import resolve_svn_profile_id
from sg_preflight.delivery_readiness import classify_changelog_text
from sg_preflight.profiles import get_run_profile
from sg_preflight.workbook_finder import (
    SOURCE_AUTO_GENERATED_LOCALLY,
    WORKBOOK_FORMAT_A_DATE_STAMPED,
    WorkbookCandidate,
    _candidate_profile_ids,
    _safe_resolve,
    _stat_workbook,
)


VALEO_COMPRESSION_FACTOR = 0.88

FORMAT_A_HEADER = (
    "Variant",
    "TextureCube",
    "Texture2D",
    "ArrayResource",
    "Effect",
    "Total",
    "Valeo est.",
)

# Per `[[feedback-team-wording]]` — natural prose, no overclaim, no codenames.
WORKBOOK_BANNER = (
    "SGFX auto-generated locally from raw BMW export-size data. "
    "Manual review remains required. Decision: not approval — evidence only."
)
OFFICIAL_DELIVERY_HEADER = (
    "Date",
    "SVN",
    "Changelog",
    "Ramses Size",
    "Logic Size",
    "Screenshot",
    "Comment",
)
OFFICIAL_WORKBOOK_COMMENT = (
    "Local preflight workbook generated from BMW export stdout. "
    "Not the official delivery record; update the shared workbook via the official delivery checklist tool."
)
_EXPORT_SIZE_RE = re.compile(
    r"File\s+sizes\s*:\s*Ramses\s*[:=]\s*(?P<ramses>\d+)\s*b\s+R\s*Logic\s*[:=]\s*(?P<logic>\d+)\s*b",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class RawVariantRow:
    variant: str
    texture_cube: float
    texture_2d: float
    array_resource: float
    effect: float

    @property
    def total(self) -> float:
        return self.texture_cube + self.texture_2d + self.array_resource + self.effect

    @property
    def valeo_estimate(self) -> float:
        return round(self.total * VALEO_COMPRESSION_FACTOR, 2)


@dataclass(frozen=True)
class RawExportSizeData:
    profile_id: str
    source_path: Path
    rows: tuple[RawVariantRow, ...]


@dataclass(frozen=True)
class ExportSizeSummary:
    ramses_size_bytes: int
    logic_size_bytes: int
    source_line: str


@dataclass(frozen=True)
class OfficialDeliveryWorkbookData:
    profile_id: str
    brand: str
    car_model: str
    ramses_size_bytes: int
    logic_size_bytes: int
    svn_revision: str = ""
    changelog_revision: str = ""
    generated_at: datetime | None = None


def _operator_outputs_root() -> Path:
    home = _safe_resolve(Path.home())
    if home is None:
        return Path.cwd() / "sgfx_outputs"
    return home / "sgfx_outputs"


def _brand_label(brand: str | None) -> str:
    value = str(brand or "BMW").strip() or "BMW"
    return value.upper() if value.casefold() in {"bmw", "mini"} else value


def _safe_sheet_title(value: str) -> str:
    title = re.sub(r"[\[\]:*?/\\]+", "_", str(value or "Profile").strip())[:31]
    return title or "Profile"


def official_delivery_workbook_path(
    profile_id: str,
    *,
    brand: str | None = "BMW",
    today: datetime | None = None,
) -> Path:
    del today
    profile = (profile_id or "").strip().upper() or "PROFILE"
    return _operator_outputs_root() / profile.lower() / "delivery-workbook" / f"Delivery Data - {_brand_label(brand)}.xlsx"


def parse_export_size_summary(text: str) -> ExportSizeSummary | None:
    """Parse BMW export stdout size lines such as `File sizes: Ramses: 123b RLogic: 0b`."""
    raw = str(text or "")
    match = _EXPORT_SIZE_RE.search(raw)
    if match:
        source_line = " ".join(part.strip() for part in match.group(0).splitlines() if part.strip())
        return ExportSizeSummary(
            ramses_size_bytes=int(match.group("ramses")),
            logic_size_bytes=int(match.group("logic")),
            source_line=source_line,
        )
    return None


def _run_svn_info_revision(target: Path | None) -> str:
    svn = shutil.which("svn")
    if not svn or target is None:
        return ""
    probe = target if target.is_dir() else target.parent
    if not probe.exists():
        return ""
    checkout_root = None
    for current in (probe, *probe.parents):
        if (current / ".svn").is_dir():
            checkout_root = current
            break
    if checkout_root is None:
        return ""
    try:
        completed = subprocess.run(
            [svn, "info", str(probe)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=18,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return ""
    if completed.returncode != 0:
        return ""
    for line in completed.stdout.splitlines():
        if line.startswith("Revision:"):
            value = line.split(":", 1)[1].strip()
            return f"r{value}" if value and not value.lower().startswith("r") else value
    return ""


def _read_changelog_revision(changelog_path: Path) -> str:
    if not changelog_path.is_file():
        return ""
    try:
        text = changelog_path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = changelog_path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return ""
    classification = classify_changelog_text(text)
    return classification.version or classification.header


def build_official_delivery_workbook_data(
    profile_id: str,
    *,
    ramses_size_bytes: int,
    logic_size_bytes: int,
    workspace: Path | str | None = None,
    bmw_root: Path | str | None = None,
    today: datetime | None = None,
) -> OfficialDeliveryWorkbookData:
    profile = resolve_svn_profile_id(profile_id).strip().upper() or (profile_id or "").strip().upper() or "PROFILE"
    brand = "BMW"
    model_root: Path | None = None
    try:
        run_profile = get_run_profile(profile, Path(workspace).resolve() if workspace is not None else None, bmw_root=bmw_root)
        brand = run_profile.brand or brand
        profile = resolve_svn_profile_id(run_profile.profile_id).strip().upper() or profile
        model_root = run_profile.source_project_root()
    except Exception:
        model_root = None
    changelog_revision = _read_changelog_revision(model_root / "CHANGELOG.md") if model_root is not None else ""
    svn_revision = _run_svn_info_revision(model_root or (Path(workspace).resolve() if workspace is not None else None))
    return OfficialDeliveryWorkbookData(
        profile_id=profile,
        brand=_brand_label(brand),
        car_model=profile,
        ramses_size_bytes=int(ramses_size_bytes),
        logic_size_bytes=int(logic_size_bytes),
        svn_revision=svn_revision,
        changelog_revision=changelog_revision,
        generated_at=today or datetime.now(timezone.utc),
    )


def generate_official_delivery_workbook(
    data: OfficialDeliveryWorkbookData,
    *,
    output_path: Path | None = None,
) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError:
        raise

    target = output_path or official_delivery_workbook_path(data.profile_id, brand=data.brand)
    target.parent.mkdir(parents=True, exist_ok=True)
    generated_at = data.generated_at or datetime.now(timezone.utc)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_sheet_title(data.car_model)
    sheet.append(list(OFFICIAL_DELIVERY_HEADER))
    sheet.append(
        [
            generated_at.strftime("%d.%m.%Y"),
            data.svn_revision,
            data.changelog_revision,
            f"{int(data.ramses_size_bytes)}b",
            f"{int(data.logic_size_bytes)}b",
            "",
            OFFICIAL_WORKBOOK_COMMENT,
        ]
    )
    workbook.save(target)
    return target


def generate_official_delivery_workbook_from_export_log(
    profile_id: str,
    export_stdout: str,
    *,
    workspace: Path | str | None = None,
    bmw_root: Path | str | None = None,
    today: datetime | None = None,
) -> WorkbookCandidate | None:
    summary = parse_export_size_summary(export_stdout)
    if summary is None:
        return None
    data = build_official_delivery_workbook_data(
        profile_id,
        ramses_size_bytes=summary.ramses_size_bytes,
        logic_size_bytes=summary.logic_size_bytes,
        workspace=workspace,
        bmw_root=bmw_root,
        today=today,
    )
    try:
        path = generate_official_delivery_workbook(data)
    except (ImportError, OSError):
        return None
    metadata = _stat_workbook(path)
    if metadata is None:
        return None
    mtime_ns, size_bytes = metadata
    return WorkbookCandidate(
        path=path,
        mtime_ns=mtime_ns,
        size_bytes=size_bytes,
        source_key="operator_local_official",
        source_classification=SOURCE_AUTO_GENERATED_LOCALLY,
        workbook_format=WORKBOOK_FORMAT_A_DATE_STAMPED,
    )


def auto_generated_workbook_path(profile_id: str, *, today: datetime | None = None) -> Path:
    """`~/sgfx_outputs/<profile_lower>/delivery-workbook/<PROFILE>_auto_<YYYYMMDD>.xlsx`"""
    profile = (profile_id or "").strip().upper() or "PROFILE"
    stamp = (today or datetime.now(timezone.utc)).strftime("%Y%m%d")
    return _operator_outputs_root() / profile.lower() / "delivery-workbook" / f"{profile}_auto_{stamp}.xlsx"


def find_raw_export_size_data(
    profile_id: str,
    *,
    workspace: Path | str | None = None,
    bmw_root: Path | str | None = None,
) -> RawExportSizeData | None:
    """Search the four documented raw-data locations + return the first parseable hit.

    Newest mtime wins when multiple candidate files exist. Supports CSV (per-variant
    rows with columns `variant,texturecube,texture2d,arrayresource,effect`) and JSON
    (list of `{variant, texturecube, ...}` dicts OR `{"variants": [...]}` wrapper).
    """
    profile = (profile_id or "").strip().upper()
    if not profile:
        return None
    ws = _safe_resolve(workspace)
    br = _safe_resolve(bmw_root)
    candidate_paths: list[Path] = []
    for variant in _candidate_profile_ids(profile):
        if br is not None:
            evo_variant = variant if variant.endswith("_EVO") else f"{variant}_EVO"
            candidate_paths.extend([
                br / "cars" / "BMW" / evo_variant / "export" / "size_data.csv",
                br / "cars" / "BMW" / evo_variant / "export" / "size_data.json",
                br / "cars" / "BMW" / variant / "export" / "size_data.csv",
                br / "cars" / "BMW" / variant / "export" / "size_data.json",
            ])
        if ws is not None:
            candidate_paths.extend([
                ws / "Cars" / "BMW" / variant / "export" / "size_data.csv",
                ws / "Cars" / "BMW" / variant / "export" / "size_data.json",
                ws / "Cars_IDCevo" / "BMW" / variant / "size_data.csv",
                ws / "Cars_IDCevo" / "BMW" / variant / "size_data.json",
            ])
    seen: set[Path] = set()
    hits: list[tuple[int, Path]] = []
    for candidate in candidate_paths:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if resolved in seen:
            continue
        seen.add(resolved)
        metadata = _stat_workbook(resolved)
        if metadata is None:
            continue
        hits.append((metadata[0], resolved))
    if not hits:
        return None
    hits.sort(key=lambda item: (item[0], str(item[1]).casefold()), reverse=True)
    for _mtime, path in hits:
        rows = _parse_raw_export_size_data(path)
        if rows:
            return RawExportSizeData(profile_id=profile, source_path=path, rows=rows)
    return None


def _parse_raw_export_size_data(path: Path) -> tuple[RawVariantRow, ...]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            return _parse_csv(path)
        if suffix == ".json":
            return _parse_json(path)
    except (OSError, ValueError, json.JSONDecodeError):
        return ()
    return ()


def _safe_float(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


_HEADER_NORMALIZE = re.compile(r"[^a-z0-9]+")


def _normalize_header(name: str) -> str:
    return _HEADER_NORMALIZE.sub("", str(name or "").casefold())


_HEADER_KEYS = {
    "variant": "variant",
    "powertrain": "variant",
    "texturecube": "texture_cube",
    "texture2d": "texture_2d",
    "arrayresource": "array_resource",
    "effect": "effect",
    "effects": "effect",
}


def _parse_csv(path: Path) -> tuple[RawVariantRow, ...]:
    rows: list[RawVariantRow] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            return ()
        canonical: dict[str, str] = {}
        for field_name in reader.fieldnames:
            mapped = _HEADER_KEYS.get(_normalize_header(field_name))
            if mapped:
                canonical[mapped] = field_name
        if "variant" not in canonical:
            return ()
        for raw_row in reader:
            variant = str(raw_row.get(canonical["variant"], "")).strip()
            if not variant:
                continue
            rows.append(
                RawVariantRow(
                    variant=variant,
                    texture_cube=_safe_float(raw_row.get(canonical.get("texture_cube", ""), 0)),
                    texture_2d=_safe_float(raw_row.get(canonical.get("texture_2d", ""), 0)),
                    array_resource=_safe_float(raw_row.get(canonical.get("array_resource", ""), 0)),
                    effect=_safe_float(raw_row.get(canonical.get("effect", ""), 0)),
                )
            )
    return tuple(rows)


def _parse_json(path: Path) -> tuple[RawVariantRow, ...]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    items: Iterable[object]
    if isinstance(raw, dict):
        items = raw.get("variants", []) if isinstance(raw.get("variants"), list) else []
    elif isinstance(raw, list):
        items = raw
    else:
        return ()
    rows: list[RawVariantRow] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        canonical: dict[str, object] = {}
        for key, value in item.items():
            mapped = _HEADER_KEYS.get(_normalize_header(str(key)))
            if mapped:
                canonical[mapped] = value
        variant = str(canonical.get("variant", "")).strip()
        if not variant:
            continue
        rows.append(
            RawVariantRow(
                variant=variant,
                texture_cube=_safe_float(canonical.get("texture_cube", 0)),
                texture_2d=_safe_float(canonical.get("texture_2d", 0)),
                array_resource=_safe_float(canonical.get("array_resource", 0)),
                effect=_safe_float(canonical.get("effect", 0)),
            )
        )
    return tuple(rows)


def generate_workbook_from_raw(
    data: RawExportSizeData,
    *,
    output_path: Path | None = None,
    today: datetime | None = None,
) -> Path:
    """Write a Format A workbook to `output_path` (default `auto_generated_workbook_path(...)`).

    Returns the resolved output path on success. Raises `ImportError` if openpyxl
    is unavailable, `OSError` if the operator-local output dir can't be created.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        raise

    target = output_path or auto_generated_workbook_path(data.profile_id, today=today)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append([data.profile_id, "", "", "", "", "", ""])
    sheet.append(list(FORMAT_A_HEADER))
    for row in data.rows:
        sheet.append([
            row.variant,
            round(row.texture_cube, 2),
            round(row.texture_2d, 2),
            round(row.array_resource, 2),
            round(row.effect, 2),
            round(row.total, 2),
            row.valeo_estimate,
        ])
    # Append the banner one row below the data so it doesn't interfere with the
    # canonical header row parsing.
    sheet.append([])
    sheet.append([WORKBOOK_BANNER])
    # Drop a sibling metadata sheet so future readers can tell the file apart
    # from a CI-produced workbook without name-matching alone.
    meta_sheet = workbook.create_sheet("SGFX Provenance")
    meta_sheet.append(["Field", "Value"])
    meta_sheet.append(["source_classification", SOURCE_AUTO_GENERATED_LOCALLY])
    meta_sheet.append(["workbook_format", WORKBOOK_FORMAT_A_DATE_STAMPED])
    meta_sheet.append(["profile_id", data.profile_id])
    meta_sheet.append(["raw_data_source", str(data.source_path)])
    meta_sheet.append([
        "generated_at_utc",
        (today or datetime.now(timezone.utc)).isoformat(timespec="seconds"),
    ])
    meta_sheet.append(["variant_count", len(data.rows)])
    meta_sheet.append(["banner", WORKBOOK_BANNER])
    workbook.save(target)
    return target


def auto_generate_if_raw_available(
    profile_id: str,
    *,
    workspace: Path | str | None = None,
    bmw_root: Path | str | None = None,
    today: datetime | None = None,
) -> WorkbookCandidate | None:
    """Top-level helper: find raw data → render Format A workbook → return candidate.

    Returns None if no raw data was found or openpyxl is unavailable. Callers
    treat None as "fall back to existing honest unavailable wording".
    """
    raw = find_raw_export_size_data(profile_id, workspace=workspace, bmw_root=bmw_root)
    if raw is None or not raw.rows:
        return None
    target = auto_generated_workbook_path(raw.profile_id, today=today)
    try:
        path = generate_workbook_from_raw(raw, output_path=target, today=today)
    except (ImportError, OSError):
        return None
    metadata = _stat_workbook(path)
    if metadata is None:
        return None
    mtime_ns, size_bytes = metadata
    return WorkbookCandidate(
        path=path,
        mtime_ns=mtime_ns,
        size_bytes=size_bytes,
        source_key="operator_local_auto_gen",
        source_classification=SOURCE_AUTO_GENERATED_LOCALLY,
        workbook_format=WORKBOOK_FORMAT_A_DATE_STAMPED,
    )
