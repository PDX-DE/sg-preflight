from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from sg_preflight.profiles import resolve_source_repo_root


EVIDENCE_ONLY_BANNER = (
    "Evidence only - export-size trends are read from local size_analysis workbooks; "
    "manual review remains required."
)
SIGNIFICANT_CHANGE_LABEL = "significant size change vs prior workbook - review"
UNREADABLE_LAYOUT_LABEL = "could not read Overview layout - review"
DETAIL_FALLBACK_LABEL = "total derived from detail sheets - review"

SIZE_ANALYSIS_RELATIVE = Path("Cars") / "size_analysis"
OVERVIEW_SHEET = "Overview"
SIGNIFICANT_CHANGE_MIN_PERCENT = 25.0
SIGNIFICANT_CHANGE_MIN_UNITS = 1000.0

_FILENAME_RE = re.compile(r"^(?P<profile>.+)_(?P<suffix>[^_]+)$")
_DATE_FILENAME_RE = re.compile(r"^\d{8}$")
_VERSION_NUMBER_RE = re.compile(r"^v(?P<number>\d+)$", re.IGNORECASE)
_NORMALIZE_RE = re.compile(r"[^a-z0-9]+")
_COLUMN_IGNORED_HEADERS = {"min", "max"}
_COLUMN_TITLE_PREFIX = "force recalculate"
_ROW_TABLE_CATEGORIES = ("TextureCube", "Texture2D", "ArrayResource", "Effect")
_COLUMN_DETAIL_CATEGORIES = ("Txt", "Mesh")


@dataclass(frozen=True)
class ExportSizeVariant:
    name: str
    date_text: str
    total: float | None
    total_source: str
    metrics: dict[str, float]
    raw_values: dict[str, str]

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "date_text": self.date_text,
            "total": self.total,
            "total_source": self.total_source,
            "metrics": dict(self.metrics),
            "raw_values": dict(self.raw_values),
        }


@dataclass(frozen=True)
class ExportSizeWorkbook:
    profile_id: str
    filename_suffix: str
    suffix_kind: str
    workbook_path: str
    relative_path: str
    title: str
    layout: str
    status: str
    semantic_date: str
    semantic_date_source: str
    version_order: int
    row_count: int
    column_count: int
    sheet_count: int
    variant_count: int
    total_min: float | None
    total_max: float | None
    total_avg: float | None
    detail_fallback_count: int
    review_flags: tuple[str, ...]
    variants: tuple[ExportSizeVariant, ...]

    @property
    def parsed(self) -> bool:
        return self.status == "parsed"

    @property
    def needs_review(self) -> bool:
        return bool(self.review_flags)

    def to_dict(self) -> dict[str, Any]:
        return {
            "profile_id": self.profile_id,
            "filename_suffix": self.filename_suffix,
            "suffix_kind": self.suffix_kind,
            "workbook_path": self.workbook_path,
            "relative_path": self.relative_path,
            "title": self.title,
            "layout": self.layout,
            "status": self.status,
            "semantic_date": self.semantic_date,
            "semantic_date_source": self.semantic_date_source,
            "version_order": self.version_order,
            "row_count": self.row_count,
            "column_count": self.column_count,
            "sheet_count": self.sheet_count,
            "variant_count": self.variant_count,
            "total_min": self.total_min,
            "total_max": self.total_max,
            "total_avg": self.total_avg,
            "detail_fallback_count": self.detail_fallback_count,
            "review_flags": list(self.review_flags),
            "needs_review": self.needs_review,
            "variants": [variant.to_dict() for variant in self.variants],
        }


@dataclass(frozen=True)
class ExportSizeTrendChange:
    profile_id: str
    previous_workbook: str
    current_workbook: str
    previous_date: str
    current_date: str
    previous_total_max: float
    current_total_max: float
    delta_total: float
    delta_percent: float
    review_flags: tuple[str, ...]

    @property
    def needs_review(self) -> bool:
        return bool(self.review_flags)

    def to_dict(self) -> dict[str, Any]:
        return {
            "profile_id": self.profile_id,
            "previous_workbook": self.previous_workbook,
            "current_workbook": self.current_workbook,
            "previous_date": self.previous_date,
            "current_date": self.current_date,
            "previous_total_max": self.previous_total_max,
            "current_total_max": self.current_total_max,
            "delta_total": self.delta_total,
            "delta_percent": self.delta_percent,
            "review_flags": list(self.review_flags),
            "needs_review": self.needs_review,
        }


@dataclass(frozen=True)
class ExportSizeTrendBoard:
    repo_root: Path
    source_state: str
    generated_at_utc: str
    workbook_dir: str
    workbooks: tuple[ExportSizeWorkbook, ...]
    trend_changes: tuple[ExportSizeTrendChange, ...]
    manual_review_banner: str = EVIDENCE_ONLY_BANNER

    @property
    def counts(self) -> dict[str, Any]:
        layout_counts = Counter(item.layout for item in self.workbooks)
        date_source_counts = Counter(item.semantic_date_source for item in self.workbooks if item.semantic_date_source)
        suffix_kind_counts = Counter(item.suffix_kind for item in self.workbooks)
        profiles = {item.profile_id for item in self.workbooks}
        parsed = [item for item in self.workbooks if item.parsed]
        unreadable = [item for item in self.workbooks if not item.parsed]
        review_workbooks = [item for item in self.workbooks if item.review_flags]
        review_changes = [item for item in self.trend_changes if item.review_flags]
        return {
            "workbook_count": len(self.workbooks),
            "profile_count": len(profiles),
            "parsed_workbook_count": len(parsed),
            "unreadable_workbook_count": len(unreadable),
            "variant_total": sum(item.variant_count for item in parsed),
            "layout_counts": dict(sorted(layout_counts.items())),
            "date_source_counts": dict(sorted(date_source_counts.items())),
            "suffix_kind_counts": dict(sorted(suffix_kind_counts.items())),
            "trend_profile_count": len({item.profile_id for item in self.trend_changes}),
            "trend_change_count": len(self.trend_changes),
            "review_change_count": len(review_changes),
            "review_workbook_count": len(review_workbooks),
            "review_profile_count": len({item.profile_id for item in review_changes}),
            "detail_fallback_count": sum(item.detail_fallback_count for item in parsed),
        }

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": 1,
            "repo_root": str(self.repo_root),
            "source_state": self.source_state,
            "generated_at_utc": self.generated_at_utc,
            "workbook_dir": self.workbook_dir,
            "manual_review_banner": self.manual_review_banner,
            "significant_change_label": SIGNIFICANT_CHANGE_LABEL,
            "unreadable_layout_label": UNREADABLE_LAYOUT_LABEL,
            "detail_fallback_label": DETAIL_FALLBACK_LABEL,
            "thresholds": {
                "significant_change_min_percent": SIGNIFICANT_CHANGE_MIN_PERCENT,
                "significant_change_min_units": SIGNIFICANT_CHANGE_MIN_UNITS,
            },
            "counts": self.counts,
            "workbooks": [item.to_dict() for item in self.workbooks],
            "trend_changes": [item.to_dict() for item in self.trend_changes],
        }


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, float):
        return format(value, ".12g")
    return str(value).strip()


def _normalize(value: str) -> str:
    return _NORMALIZE_RE.sub("", value.casefold())


def _numeric(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _filename_parts(path: Path) -> tuple[str, str, str, int]:
    match = _FILENAME_RE.match(path.stem)
    if match is None:
        return path.stem.upper(), "", "unknown", 0
    profile = match.group("profile").upper()
    suffix = match.group("suffix")
    if _DATE_FILENAME_RE.fullmatch(suffix):
        return profile, suffix, "filename_date", 0
    version_match = _VERSION_NUMBER_RE.fullmatch(suffix)
    if version_match is not None:
        return profile, suffix, "version_number", int(version_match.group("number"))
    if suffix.casefold() == "vx":
        return profile, suffix, "version_vx", 10_000
    return profile, suffix, "unknown", 0


def _filename_date_text(suffix: str) -> str:
    if not _DATE_FILENAME_RE.fullmatch(suffix):
        return ""
    return f"{suffix[0:4]}-{suffix[4:6]}-{suffix[6:8]}"


def _parse_date_text(value: str) -> datetime | None:
    text = str(value).strip()
    for fmt in ("%d.%m.%y %H.%M", "%d.%m.%Y %H.%M", "%d.%m.%y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    if _DATE_FILENAME_RE.fullmatch(text):
        try:
            return datetime.strptime(text, "%Y%m%d")
        except ValueError:
            return None
    return None


def _date_sort_tuple(value: str, suffix: str, version_order: int) -> tuple[datetime, int, str]:
    parsed = _parse_date_text(value)
    if parsed is None:
        parsed = _parse_date_text(suffix)
    return parsed or datetime(1900, 1, 1), version_order, suffix.casefold()


def _semantic_date_from_sheet(values: list[str]) -> str:
    parsed_values = [parsed for parsed in (_parse_date_text(value) for value in values) if parsed is not None]
    if parsed_values:
        return max(parsed_values).isoformat(sep=" ", timespec="minutes")
    return max(values) if values else ""


def _source_repo(repo_root: Path | None, workspace_root: Path | None) -> Path:
    if repo_root is not None:
        return Path(repo_root).resolve()
    return resolve_source_repo_root(workspace_root)


def _relative_to_repo(path: Path, repo_root: Path) -> str:
    try:
        return str(path.relative_to(repo_root)).replace("\\", "/")
    except ValueError:
        return str(path)


def _rows_from_sheet(sheet: Any) -> list[list[object]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _first_title(rows: list[list[object]], fallback: str) -> str:
    for row in rows:
        if not row:
            continue
        text = _cell_text(row[0])
        if text and _normalize(text) != "variant":
            return text
    return fallback


def _find_row_table_header(rows: list[list[object]]) -> tuple[int, list[str]]:
    for index, row in enumerate(rows):
        headers = [_cell_text(value) for value in row]
        normalized = [_normalize(value) for value in headers]
        if normalized and normalized[0] == "variant" and "total" in normalized:
            return index, headers
    return -1, []


def _find_column_header(rows: list[list[object]]) -> tuple[int, list[str]]:
    for index, row in enumerate(rows[:10]):
        values = [_cell_text(value) for value in row]
        if not values or values[0]:
            continue
        candidates = [
            value
            for value in values[1:]
            if value
            and value.casefold() not in _COLUMN_IGNORED_HEADERS
            and not value.casefold().startswith(_COLUMN_TITLE_PREFIX)
        ]
        if candidates:
            return index, values
    return -1, []


def _row_by_label(rows: list[list[object]]) -> dict[str, list[object]]:
    labels: dict[str, list[object]] = {}
    for row in rows:
        if not row:
            continue
        label = _normalize(_cell_text(row[0]))
        if label:
            labels[label] = row
    return labels


def _sheet_sum(loaded: Any, sheet_name: str) -> float | None:
    if sheet_name not in loaded.sheetnames:
        return None
    sheet = loaded[sheet_name]
    total = 0.0
    count = 0
    for row in sheet.iter_rows(values_only=True):
        if len(row) < 2:
            continue
        value = _numeric(row[1])
        if value is not None:
            total += value
            count += 1
    return total if count else None


def _detail_total(loaded: Any, variant_name: str, layout: str) -> float | None:
    suffixes = _COLUMN_DETAIL_CATEGORIES if layout == "column_matrix" else _ROW_TABLE_CATEGORIES
    total = 0.0
    count = 0
    for suffix in suffixes:
        value = _sheet_sum(loaded, f"{variant_name} {suffix}")
        if value is None:
            continue
        total += value
        count += 1
    return total if count else None


def _variant_from_values(
    *,
    loaded: Any,
    name: str,
    date_text: str,
    total_value: object,
    metrics: dict[str, object],
    layout: str,
) -> ExportSizeVariant:
    numeric_total = _numeric(total_value)
    total_source = "overview"
    review_metrics: dict[str, float] = {}
    raw_values: dict[str, str] = {}
    for key, value in metrics.items():
        raw_values[key] = _cell_text(value)
        numeric_value = _numeric(value)
        if numeric_value is not None:
            review_metrics[key] = numeric_value
    if numeric_total is None:
        detail_value = _detail_total(loaded, name, layout)
        if detail_value is not None:
            numeric_total = detail_value
            total_source = "detail_sheets"
    raw_values["Total"] = _cell_text(total_value)
    if date_text:
        raw_values["Date"] = date_text
    return ExportSizeVariant(
        name=name,
        date_text=date_text,
        total=numeric_total,
        total_source=total_source,
        metrics=review_metrics,
        raw_values=raw_values,
    )


def _parse_row_table(loaded: Any, rows: list[list[object]]) -> tuple[tuple[ExportSizeVariant, ...], str, str, int]:
    header_index, headers = _find_row_table_header(rows)
    if header_index < 0:
        return (), "", "", 0
    variants: list[ExportSizeVariant] = []
    total_index = next((index for index, header in enumerate(headers) if _normalize(header) == "total"), -1)
    for row in rows[header_index + 1 :]:
        name = _cell_text(row[0] if row else "")
        if not name:
            continue
        metrics: dict[str, object] = {}
        for index, header in enumerate(headers[1:], start=1):
            label = header.strip()
            if not label or index >= len(row):
                continue
            metrics[label] = row[index]
        total_value = row[total_index] if 0 <= total_index < len(row) else None
        if _numeric(total_value) is None and not any(_numeric(value) is not None for value in metrics.values()):
            continue
        variants.append(
            _variant_from_values(
                loaded=loaded,
                name=name,
                date_text="",
                total_value=total_value,
                metrics=metrics,
                layout="row_table",
            )
        )
    return tuple(variants), "row_table", "Variant row table", header_index + 1


def _parse_column_matrix(loaded: Any, rows: list[list[object]]) -> tuple[tuple[ExportSizeVariant, ...], str, str, int]:
    header_index, headers = _find_column_header(rows)
    if header_index < 0:
        return (), "", "", 0
    label_rows = _row_by_label(rows)
    date_row = label_rows.get("date", [])
    total_row = label_rows.get("total", [])
    metric_rows = {
        "Textures": label_rows.get("textures", []),
        "Meshes": label_rows.get("meshes", []),
        "Valeo est.": label_rows.get("valeoest", []),
        "Valeo": label_rows.get("valeo", []),
        "Ratio": label_rows.get("ratio", []),
    }
    variants: list[ExportSizeVariant] = []
    for index, header in enumerate(headers[1:], start=1):
        name = header.strip()
        if not name or name.casefold() in _COLUMN_IGNORED_HEADERS:
            continue
        date_text = _cell_text(date_row[index] if index < len(date_row) else "")
        total_value = total_row[index] if index < len(total_row) else None
        metrics = {
            label: row[index]
            for label, row in metric_rows.items()
            if row and index < len(row) and _cell_text(row[index])
        }
        if _numeric(total_value) is None and not any(_numeric(value) is not None for value in metrics.values()):
            continue
        variants.append(
            _variant_from_values(
                loaded=loaded,
                name=name,
                date_text=date_text,
                total_value=total_value,
                metrics=metrics,
                layout="column_matrix",
            )
        )
    return tuple(variants), "column_matrix", "Variant column matrix", header_index + 1


def _workbook_from_unreadable(
    *,
    repo_root: Path,
    path: Path,
    profile_id: str,
    suffix: str,
    suffix_kind: str,
    version_order: int,
    status: str,
    reason: str,
) -> ExportSizeWorkbook:
    date_text = _filename_date_text(suffix)
    return ExportSizeWorkbook(
        profile_id=profile_id,
        filename_suffix=suffix,
        suffix_kind=suffix_kind,
        workbook_path=str(path),
        relative_path=_relative_to_repo(path, repo_root),
        title=path.stem,
        layout="unreadable",
        status=status,
        semantic_date=date_text,
        semantic_date_source="filename" if date_text else "",
        version_order=version_order,
        row_count=0,
        column_count=0,
        sheet_count=0,
        variant_count=0,
        total_min=None,
        total_max=None,
        total_avg=None,
        detail_fallback_count=0,
        review_flags=(f"{UNREADABLE_LAYOUT_LABEL}: {reason}",),
        variants=(),
    )


def _parse_workbook(repo_root: Path, path: Path) -> ExportSizeWorkbook:
    profile_id, suffix, suffix_kind, version_order = _filename_parts(path)
    try:
        loaded = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return _workbook_from_unreadable(
            repo_root=repo_root,
            path=path,
            profile_id=profile_id,
            suffix=suffix,
            suffix_kind=suffix_kind,
            version_order=version_order,
            status="unreadable",
            reason=str(exc),
        )
    try:
        if OVERVIEW_SHEET not in loaded.sheetnames:
            return _workbook_from_unreadable(
                repo_root=repo_root,
                path=path,
                profile_id=profile_id,
                suffix=suffix,
                suffix_kind=suffix_kind,
                version_order=version_order,
                status="no_overview_sheet",
                reason="Overview sheet missing",
            )
        worksheet = loaded[OVERVIEW_SHEET]
        rows = _rows_from_sheet(worksheet)
        title = _first_title(rows, path.stem)
        variants, layout, _layout_note, _header_row = _parse_row_table(loaded, rows)
        if not variants:
            variants, layout, _layout_note, _header_row = _parse_column_matrix(loaded, rows)
        if not variants:
            return ExportSizeWorkbook(
                profile_id=profile_id,
                filename_suffix=suffix,
                suffix_kind=suffix_kind,
                workbook_path=str(path),
                relative_path=_relative_to_repo(path, repo_root),
                title=title,
                layout="unreadable",
                status="unreadable_layout",
                semantic_date=_filename_date_text(suffix),
                semantic_date_source="filename" if _filename_date_text(suffix) else "",
                version_order=version_order,
                row_count=int(worksheet.max_row or 0),
                column_count=int(worksheet.max_column or 0),
                sheet_count=len(loaded.sheetnames),
                variant_count=0,
                total_min=None,
                total_max=None,
                total_avg=None,
                detail_fallback_count=0,
                review_flags=(UNREADABLE_LAYOUT_LABEL,),
                variants=(),
            )
        filename_date = _filename_date_text(suffix)
        sheet_dates = [variant.date_text for variant in variants if variant.date_text]
        semantic_date = filename_date or _semantic_date_from_sheet(sheet_dates)
        semantic_source = "filename" if filename_date else ("overview_date_row" if sheet_dates else "")
        totals = [variant.total for variant in variants if variant.total is not None]
        detail_count = sum(1 for variant in variants if variant.total_source == "detail_sheets")
        flags: list[str] = []
        if detail_count:
            flags.append(DETAIL_FALLBACK_LABEL)
        return ExportSizeWorkbook(
            profile_id=profile_id,
            filename_suffix=suffix,
            suffix_kind=suffix_kind,
            workbook_path=str(path),
            relative_path=_relative_to_repo(path, repo_root),
            title=title,
            layout=layout,
            status="parsed",
            semantic_date=semantic_date,
            semantic_date_source=semantic_source,
            version_order=version_order,
            row_count=int(worksheet.max_row or 0),
            column_count=int(worksheet.max_column or 0),
            sheet_count=len(loaded.sheetnames),
            variant_count=len(variants),
            total_min=min(totals) if totals else None,
            total_max=max(totals) if totals else None,
            total_avg=(sum(totals) / len(totals)) if totals else None,
            detail_fallback_count=detail_count,
            review_flags=tuple(flags),
            variants=variants,
        )
    finally:
        loaded.close()


def _sort_key(workbook: ExportSizeWorkbook) -> tuple[datetime, int, str]:
    return _date_sort_tuple(workbook.semantic_date, workbook.filename_suffix, workbook.version_order)


def _trend_changes(workbooks: tuple[ExportSizeWorkbook, ...]) -> tuple[ExportSizeTrendChange, ...]:
    by_profile: dict[str, list[ExportSizeWorkbook]] = defaultdict(list)
    for workbook in workbooks:
        if workbook.parsed and workbook.total_max is not None:
            by_profile[workbook.profile_id].append(workbook)
    changes: list[ExportSizeTrendChange] = []
    for profile_id, profile_workbooks in sorted(by_profile.items()):
        ordered = sorted(profile_workbooks, key=_sort_key)
        for previous, current in zip(ordered, ordered[1:]):
            if previous.total_max is None or current.total_max is None or previous.total_max == 0:
                continue
            delta = current.total_max - previous.total_max
            percent = (delta / previous.total_max) * 100.0
            flags: list[str] = []
            if abs(delta) >= SIGNIFICANT_CHANGE_MIN_UNITS and abs(percent) >= SIGNIFICANT_CHANGE_MIN_PERCENT:
                flags.append(SIGNIFICANT_CHANGE_LABEL)
            changes.append(
                ExportSizeTrendChange(
                    profile_id=profile_id,
                    previous_workbook=previous.relative_path,
                    current_workbook=current.relative_path,
                    previous_date=previous.semantic_date,
                    current_date=current.semantic_date,
                    previous_total_max=previous.total_max,
                    current_total_max=current.total_max,
                    delta_total=delta,
                    delta_percent=percent,
                    review_flags=tuple(flags),
                )
            )
    return tuple(changes)


def build_export_size_trend_board(
    repo_root: Path | None = None,
    *,
    workspace_root: Path | None = None,
) -> ExportSizeTrendBoard:
    source_root = _source_repo(repo_root, workspace_root)
    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    workbook_dir = source_root / SIZE_ANALYSIS_RELATIVE
    if not workbook_dir.is_dir():
        return ExportSizeTrendBoard(
            repo_root=source_root,
            source_state="missing_size_analysis",
            generated_at_utc=generated_at,
            workbook_dir=str(workbook_dir),
            workbooks=(),
            trend_changes=(),
        )
    workbooks = tuple(_parse_workbook(source_root, path) for path in sorted(workbook_dir.glob("*.xlsx")))
    return ExportSizeTrendBoard(
        repo_root=source_root,
        source_state="ready",
        generated_at_utc=generated_at,
        workbook_dir=str(workbook_dir),
        workbooks=tuple(sorted(workbooks, key=lambda item: (item.profile_id, _sort_key(item)))),
        trend_changes=_trend_changes(workbooks),
    )


def export_size_trend_markdown(board: ExportSizeTrendBoard) -> str:
    payload = board.to_dict()
    counts = payload["counts"]
    lines = [
        "# Export Size Trend",
        "",
        EVIDENCE_ONLY_BANNER,
        "",
        f"- source: `{payload['repo_root']}`",
        f"- workbook folder: `{payload['workbook_dir']}`",
        f"- generated: `{payload['generated_at_utc']}`",
        f"- workbooks: {counts['workbook_count']}",
        f"- profiles: {counts['profile_count']}",
        f"- parsed workbooks: {counts['parsed_workbook_count']}",
        f"- trend changes: {counts['trend_change_count']}",
        f"- review changes: {counts['review_change_count']}",
        "",
        "| Profile | Previous | Current | Previous max Total | Current max Total | Delta | Percent | Review flags |",
        "| --- | --- | --- | ---: | ---: | ---: | ---: | --- |",
    ]
    for change in board.trend_changes:
        lines.append(
            "| "
            + " | ".join(
                (
                    change.profile_id,
                    Path(change.previous_workbook).name,
                    Path(change.current_workbook).name,
                    _format_number(change.previous_total_max),
                    _format_number(change.current_total_max),
                    _format_number(change.delta_total),
                    f"{change.delta_percent:.2f}%",
                    "; ".join(change.review_flags).replace("|", "\\|"),
                )
            )
            + " |"
        )
    if any(not item.parsed for item in board.workbooks):
        lines.extend(("", "## Workbook Review", ""))
        for workbook in board.workbooks:
            if workbook.parsed:
                continue
            lines.append(f"- `{workbook.relative_path}`: {'; '.join(workbook.review_flags)}")
    return "\n".join(lines).rstrip() + "\n"


def write_export_size_trend_board(board: ExportSizeTrendBoard, output_root: Path) -> dict[str, str]:
    output_root.mkdir(parents=True, exist_ok=True)
    json_path = output_root / "export-size-trend.json"
    markdown_path = output_root / "export-size-trend.md"
    json_path.write_text(json.dumps(board.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    markdown_path.write_text(export_size_trend_markdown(board), encoding="utf-8")
    return {
        "json_path": str(json_path),
        "markdown_path": str(markdown_path),
    }


def _format_number(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.2f}"
