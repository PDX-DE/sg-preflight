from __future__ import annotations

import io
import json
from contextlib import redirect_stdout
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from sg_preflight.cli import main
from sg_preflight.export_size_trend import (
    DETAIL_FALLBACK_LABEL,
    SIGNIFICANT_CHANGE_LABEL,
    UNREADABLE_LAYOUT_LABEL,
    build_export_size_trend_board,
)


def _write_column_matrix(
    path: Path,
    *,
    title: str,
    variant_totals: dict[str, float | None],
    date_text: str,
    detail_totals: dict[str, tuple[float, float]] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    variants = list(variant_totals)
    sheet.append([title, "", "Force recalculate with Ctrl-Alt-F9"])
    sheet.append([])
    sheet.append(["", *variants, "", "Min", "Max"])
    sheet.append(["Date", *[date_text for _ in variants]])
    sheet.append(["Total", *[variant_totals[name] for name in variants]])
    sheet.append(["Textures", *[9000 for _ in variants]])
    sheet.append(["Meshes", *[3000 for _ in variants]])
    sheet.append([])
    sheet.append(["Valeo est.", *[10000 for _ in variants]])
    for variant, (txt_total, mesh_total) in (detail_totals or {}).items():
        txt = workbook.create_sheet(f"{variant} Txt")
        txt.append(["texture-a", txt_total])
        mesh = workbook.create_sheet(f"{variant} Mesh")
        mesh.append(["mesh-a", mesh_total])
    workbook.save(path)


def _write_row_table(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append(["G65"])
    sheet.append(["Variant", "TextureCube", "Texture2D", "ArrayResource", "Effect", "Total", "Valeo est."])
    sheet.append(["BEV-Basis", 5616, 11935.61, 10677.94, 435.56, 28665.11, 25225.3])
    sheet.append(["ICE-Basis", 5616, 11935.61, 10677.94, 435.56, 28515.26, 25225.3])
    workbook.save(path)


def _write_unreadable(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append(["BAD Sizes - vx"])
    sheet.append(["nothing useful here"])
    workbook.save(path)


def _fixture(root: Path) -> Path:
    repo = root / "repositories" / "trunk"
    size_root = repo / "Cars" / "size_analysis"
    _write_column_matrix(
        size_root / "F55_v8.xlsx",
        title="F55 Sizes - v8",
        variant_totals={"ICE-Basis": 111207.38},
        date_text="10.02.23 11.15",
    )
    _write_column_matrix(
        size_root / "F55_v9.xlsx",
        title="F55 Sizes - vx",
        variant_totals={"ICE-Basis": 20984.53},
        date_text="05.10.23 08.30",
    )
    _write_column_matrix(
        size_root / "F55_vx.xlsx",
        title="F55 Sizes - vx",
        variant_totals={"ICE-Basis": 20847.19},
        date_text="28.08.24 15.11",
    )
    _write_row_table(size_root / "G65_20251002.xlsx")
    _write_column_matrix(
        size_root / "DET_v1.xlsx",
        title="DET Sizes - v1",
        variant_totals={"ICE-Basis": None},
        date_text="01.01.24 10.00",
        detail_totals={"ICE-Basis": (60.0, 40.0)},
    )
    _write_unreadable(size_root / "BAD_vx.xlsx")
    return repo


class TestExportSizeTrend(unittest.TestCase):
    def test_board_parses_layouts_dates_and_cautious_review_flags(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            repo = _fixture(root)

            board = build_export_size_trend_board(repo, workspace_root=root)

        payload = board.to_dict()
        self.assertEqual(payload["source_state"], "ready")
        self.assertEqual(payload["counts"]["workbook_count"], 6)
        self.assertEqual(payload["counts"]["profile_count"], 4)
        self.assertEqual(payload["counts"]["parsed_workbook_count"], 5)
        self.assertEqual(payload["counts"]["unreadable_workbook_count"], 1)
        self.assertEqual(payload["counts"]["layout_counts"]["column_matrix"], 4)
        self.assertEqual(payload["counts"]["layout_counts"]["row_table"], 1)
        self.assertEqual(payload["counts"]["layout_counts"]["unreadable"], 1)
        self.assertEqual(payload["counts"]["trend_change_count"], 2)
        self.assertEqual(payload["counts"]["review_change_count"], 1)
        self.assertNotIn("failed", json.dumps(payload).lower())

        f55_v9 = next(item for item in payload["workbooks"] if item["relative_path"].endswith("F55_v9.xlsx"))
        self.assertEqual(f55_v9["semantic_date"], "2023-10-05 08:30")
        self.assertEqual(f55_v9["semantic_date_source"], "overview_date_row")

        g65 = next(item for item in payload["workbooks"] if item["profile_id"] == "G65")
        self.assertEqual(g65["layout"], "row_table")
        self.assertEqual(g65["semantic_date"], "2025-10-02")
        self.assertEqual(g65["semantic_date_source"], "filename")
        self.assertEqual(g65["variant_count"], 2)

        det = next(item for item in payload["workbooks"] if item["profile_id"] == "DET")
        self.assertIn(DETAIL_FALLBACK_LABEL, det["review_flags"])
        self.assertEqual(det["total_max"], 100.0)

        bad = next(item for item in payload["workbooks"] if item["profile_id"] == "BAD")
        self.assertEqual(bad["status"], "unreadable_layout")
        self.assertIn(UNREADABLE_LAYOUT_LABEL, bad["review_flags"][0])

        review_change = next(item for item in payload["trend_changes"] if item["review_flags"])
        self.assertEqual(review_change["profile_id"], "F55")
        self.assertIn("F55_v8.xlsx", review_change["previous_workbook"])
        self.assertIn("F55_v9.xlsx", review_change["current_workbook"])
        self.assertLess(review_change["delta_percent"], -79.0)
        self.assertIn(SIGNIFICANT_CHANGE_LABEL, review_change["review_flags"])

    def test_cli_writes_export_size_trend_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            repo = _fixture(root)
            output_root = root / "out" / "export-size-trend"
            stdout = io.StringIO()

            with redirect_stdout(stdout):
                result = main(
                    [
                        "export-size-trend",
                        "--workspace",
                        str(root),
                        "--repo-root",
                        str(repo),
                        "--output-root",
                        str(output_root),
                        "--json",
                    ]
                )

            self.assertEqual(result, 0)
            payload = json.loads(stdout.getvalue())
            self.assertEqual(payload["counts"]["workbook_count"], 6)
            self.assertEqual(payload["counts"]["review_change_count"], 1)
            self.assertTrue((output_root / "export-size-trend.json").exists())
            self.assertTrue((output_root / "export-size-trend.md").exists())


if __name__ == "__main__":
    unittest.main()
